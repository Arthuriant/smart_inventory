"""Builds Bulk_Import_Template_Master.xlsx, the master copy used by the flow
"Smart Inventory - Bulk Import Template".

The master has empty list tables on the Lists sheet; the flow copies the master and fills
tblArea / tblCategory / tblRack with the current data of the requesting user's business line
(sent by the app), so the dropdowns in the downloaded template are always up to date.
Row 2 of each list table is a hidden empty row (a table needs one data row); data starts at row 3.

Usage: python build_template.py [output.xlsx]
"""
import os
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableColumn, TableFormula, TableStyleInfo

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "Bulk_Import_Template_Master.xlsx")
F = "Arial"
NAVY = "00126B"
HDR = PatternFill("solid", fgColor=NAVY)
thin = Side(style="thin", color="E2E6EF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
LAST = 500  # last list row the dropdowns look at
UOMS = ["EA", "M", "SET", "DR", "JT", "BOX", "L", "PR", "SX", "FT", "GAL", "KG", "BBL", "MT"]
HEADERS = ["Area", "Item name", "BPN", "SPN", "Category", "UOM", "Qty", "Min in area", "Rack", "Note", "Status"]
WIDTHS = [28, 30, 16, 14, 14, 9, 8, 12, 16, 30, 60]
# hidden helper columns of the Status check (inside tblImport so they fill down; the upload flow ignores them)
HELPERS = ["_Key", "_InCatalog"]
ROWS = 1000  # rows prepared with formulas; rows pasted below them get the formulas from the table

wb = Workbook()

# ---------- Import ----------
ws = wb.active
ws.title = "Import"
for c, h in enumerate(HEADERS, 1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.font = Font(name=F, bold=True, color="FFFFFF")
    cell.fill = HDR
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.column_dimensions[cell.column_letter].width = WIDTHS[c - 1]
for r in range(2, ROWS + 1):
    for c in range(1, 11):
        cell = ws.cell(row=r, column=c)
        cell.font = Font(name=F)
        cell.border = BOX
        if c in (7, 8):
            cell.number_format = "0"
for i, h in enumerate(HELPERS, len(HEADERS) + 1):
    ws.cell(row=1, column=i, value=h).font = Font(name=F, color="FFFFFF")
    ws.column_dimensions[ws.cell(row=1, column=i).column_letter].hidden = True

# ---------- Status: the same checks as the app's Check (scr_bulk btnBlkCheck) ----------
# A Area, B Item name, C BPN, D SPN, E Category, F UOM, G Qty, H Min, I Rack, J Note, K Status, L _Key, M _InCatalog
# Item identity: the item name (description); SPN and BPN may repeat. The catalog names come from the app as
# text chunks "|n:<name>|n:<name>|" in Lists tblCatalog.
# The formulas use table references so they cover every row of tblImport (also rows pasted below row ROWS:
# the columns are calculated columns, Excel fills them into new rows). Text compares are exact
# (no COUNTIF/SEARCH: their * ? ~ wildcards gave wrong matches) and, like the app, ignore case.
def tr(col):
    return f"tblImport[[#This Row],[{col}]]"


def tc(col):
    return f"tblImport[{col}]"


AREA, NAME, BPN, SPN, CAT, UOM, QTY, MIN, RACK = (tr(h) for h in HEADERS[:9])
XKEY, XINC = (tr(h) for h in HELPERS)
in_cat = lambda prefix, ref: f'SUMPRODUCT(--ISNUMBER(FIND(LOWER("|{prefix}:"&{ref}&"|"),LOWER(tblCatalog[Keys]))))>0'
in_list = lambda col, ref: f'SUMPRODUCT(--({col}={ref}))>0'
FORMULAS = {
    "_Key": f'IF(OR(COUNTA(tblImport[[#This Row],[Area]:[Rack]])=0,TRIM({NAME})=""),"",TRIM({AREA})&"|"&TRIM({NAME}))',
    "_InCatalog": f'IF(TRIM({NAME})="",FALSE,{in_cat("n", f"TRIM({NAME})")})',
}
no_id = f'TRIM({NAME})=""'
new_item = f'AND(NOT({XINC}),NOT({no_id}))'
bad_num = lambda ref: f'AND({ref}<>"",IFERROR(--{ref}<0,TRUE))'
area_ok = in_list("tblArea[Area]", AREA)
# the sheet row of the first other row with the same key (the row above it if there is one, else the next one)
first = f'MATCH({XKEY},{tc("_Key")},0)+1'
dup_row = (f'IFERROR(IF({first}<>ROW(),{first},ROW()+MATCH({XKEY},'
           f'INDEX({tc("_Key")},ROW()):INDEX({tc("_Key")},ROWS({tc("_Key")})),0)),"?")')
dup_n = f'SUMPRODUCT(--({tc("_Key")}={XKEY}))'
errs = "&".join([
    f'IF({AREA}="","Area is empty. ",IF(NOT({area_ok}),"Area \'"&{AREA}&"\' is not in your business line. ",""))',
    f'IF({no_id},"Item name is needed (it identifies the item). ","")',
    f'IF({bad_num(QTY)},"Qty must be a number of 0 or more. ","")',
    f'IF({bad_num(MIN)},"Min must be a number of 0 or more. ","")',
    f'IF(AND({RACK}<>"",{AREA}<>"",{area_ok},SUMPRODUCT(--(tblRack[Rack]={RACK}),--(tblRack[Rack is in area]={AREA}))=0),'
    f'"Rack \'"&{RACK}&"\' is not in this area. ","")',
    f'IF(AND({new_item},OR({CAT}="",NOT(OR({in_list("tblCategory[Category code]", CAT)},'
    f'{in_list("tblCategory[Category name]", CAT)})))),"Category \'"&{CAT}&"\' not found (needed for a new item). ","")',
    f'IF(AND({new_item},OR({UOM}="",NOT({in_list("tblUom[UOM]", UOM)}))),'
    f'"UOM \'"&{UOM}&"\' is not valid (needed for a new item). ","")',
    f'IF(AND({XKEY}<>"",{dup_n}>1),"Same item name as row "&{dup_row}&IF({dup_n}>2," (and "&({dup_n}-2)&" more)","")&'
    f'" in this area (one rack per item per area: merge the rows, or give the item on the other rack its own name). ","")',
])
FORMULAS["Status"] = (f'IF(COUNTA(tblImport[[#This Row],[Area]:[Rack]])=0,"",'
                      f'IF(({errs})="",IF({XINC},"OK","OK - new item"),TRIM({errs})))')
COLS = HEADERS + HELPERS
for name, f in FORMULAS.items():
    c = COLS.index(name) + 1
    for r in range(2, ROWS + 1):
        cell = ws.cell(row=r, column=c, value="=" + f)
        if name == "Status":
            cell.font = Font(name=F, size=9)
            cell.border = BOX
# a row is red until its Status is OK; new items amber
ws.conditional_formatting.add(f"A2:K{ROWS}", FormulaRule(formula=['AND($K2<>"",LEFT($K2,2)<>"OK")'], stopIfTrue=True,
                                                         fill=PatternFill("solid", fgColor="FDE7E9")))
ws.conditional_formatting.add(f"K2:K{ROWS}", FormulaRule(formula=['$K2="OK - new item"'], font=Font(color="BC7600")))
ws.conditional_formatting.add(f"K2:K{ROWS}", FormulaRule(formula=['$K2="OK"'], font=Font(color="107C41")))
ws.row_dimensions[1].height = 22
ws.freeze_panes = "A2"
t = Table(displayName="tblImport", ref=f"A1:M{ROWS}")
t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
# calculated columns: Excel fills these formulas into rows added to the table (also by pasting below it)
t.tableColumns = [TableColumn(id=i, name=n, calculatedColumnFormula=TableFormula(attr_text=FORMULAS[n]) if n in FORMULAS else None)
                  for i, n in enumerate(COLS, 1)]
ws.add_table(t)

# ---------- Lists (filled by the flow) ----------
ls = wb.create_sheet("Lists")
# (table, first column, headers); one empty column between tables so they can grow independently
lists = [("tblArea", 1, ["Area"]), ("tblCategory", 3, ["Category code", "Category name"]),
         ("tblUom", 6, ["UOM"]), ("tblRack", 8, ["Rack", "Rack is in area"])]
for name, col, heads in lists:
    for i, h in enumerate(heads):
        cell = ls.cell(row=1, column=col + i, value=h)
        cell.font = Font(name=F, bold=True, color="FFFFFF")
        cell.fill = HDR
        ls.column_dimensions[cell.column_letter].width = 28 if h in ("Area", "Rack is in area") else 16
    last_row = 2 + (len(UOMS) if name == "tblUom" else 0)
    first = ls.cell(row=1, column=col).column_letter
    end = ls.cell(row=1, column=col + len(heads) - 1).column_letter
    tb = Table(displayName=name, ref=f"{first}1:{end}{last_row}")
    tb.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
    ls.add_table(tb)
for i, u in enumerate(UOMS, 3):
    ls.cell(row=i, column=6, value=u).font = Font(name=F)
# catalog keys for the Status column (filled by the flow, a few long text rows); hidden
ls.cell(row=1, column=15, value="Keys").font = Font(name=F, bold=True, color="FFFFFF")
ls.cell(row=1, column=15).fill = HDR
tc = Table(displayName="tblCatalog", ref="O1:O2")
tc.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
ls.add_table(tc)
ls.column_dimensions["O"].hidden = True
ls.row_dimensions[2].hidden = True
ls.cell(row=1, column=11, value="Filled automatically when the template is downloaded from the app. Do not edit.").font = \
    Font(name=F, italic=True, color="64707F")
# the only entry of the Rack dropdown when the row has no area yet / the area has no racks
# (column M is outside every list table, so the flow's inserts never move these cells)
HINT_NO_AREA = "M3"
HINT_NO_RACK = "M4"
ls.cell(row=3, column=13, value="(Select the Area first)").font = Font(name=F, italic=True, color="64707F")
ls.cell(row=4, column=13, value="(No rack in this area)").font = Font(name=F, italic=True, color="64707F")


def dv(formula, rng, title, msg, prompt=None):
    v = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=True,
                       errorStyle="warning", errorTitle=title, error=msg)
    if prompt:
        v.showInputMessage = True
        v.promptTitle = title
        v.prompt = prompt
    ws.add_data_validation(v)
    v.add(rng)


# Lists ranges go through INDIRECT: every row the flow adds to a Lists table inserts cells there, and Excel
# would shift a plain Lists!$A$3 reference down past the data (the dropdowns came out empty)
def lst(rng):
    return f'INDIRECT("Lists!{rng}")'


dv(f"OFFSET({lst('$A$3')},0,0,MAX(1,COUNTA({lst(f'$A$3:$A${LAST}')})),1)", f"A2:A{ROWS}", "Area",
   "This area is not in your business line. The app will mark the row red.")
dv(f"OFFSET({lst('$C$3')},0,0,MAX(1,COUNTA({lst(f'$C$3:$C${LAST}')})),1)", f"E2:E{ROWS}", "Category",
   "Use a category code from the Lists sheet (the name also works).")
dv(lst(f"$F$3:$F${len(UOMS) + 2}"), f"F2:F{ROWS}", "UOM", "UOM must be one of the values in the Lists sheet.")
# only the racks of the row's area: the block in column H that starts at the first match of the area in column I;
# a single hint entry while the row has no area yet, or when the area has no racks
dv(f"IF($A2=\"\",{lst('$' + HINT_NO_AREA[0] + '$' + HINT_NO_AREA[1:])},"
   f"IF(COUNTIF({lst(f'$I$3:$I${LAST}')},$A2)=0,{lst('$' + HINT_NO_RACK[0] + '$' + HINT_NO_RACK[1:])},"
   f"OFFSET({lst('$H$2')},MATCH($A2,{lst(f'$I$3:$I${LAST}')},0),0,COUNTIF({lst(f'$I$3:$I${LAST}')},$A2),1)))",
   f"I2:I{ROWS}", "Rack", "This rack is not in the area of this row (see the Lists sheet).",
   "Select the Area of this row first. The list shows only the racks of that area. Leave empty for no rack.")
# Rack cell greyed out while the row has no area; red when a rack (or the hint) is filled without an area
ws.conditional_formatting.add(f"I2:I{ROWS}", FormulaRule(formula=['AND($A2="",$I2<>"")'], stopIfTrue=True,
                                                     fill=PatternFill("solid", fgColor="FDE7E9"),
                                                     font=Font(color="C42B1C")))
ws.conditional_formatting.add(f"I2:I{ROWS}", FormulaRule(formula=['$A2=""'], fill=PatternFill("solid", fgColor="F1F3F6")))
num = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                     showErrorMessage=True, errorTitle="Number", error="Enter a number of 0 or more.")
ws.add_data_validation(num)
num.add(f"G2:H{ROWS}")

# ---------- Petunjuk ----------
gs = wb.create_sheet("Petunjuk", 0)
gs.column_dimensions["A"].width = 4
gs.column_dimensions["B"].width = 110
lines = [
    ("Bulk Import template - Smart Inventory", "title"),
    ("", None),
    ("Cara pakai", "h"),
    ("1. Isi sheet 'Import' mulai baris 2. Dropdown Area, Category, UOM dan Rack berisi data terkini business line Anda.", None),
    ("2. Lihat kolom Status: baris merah masih error (alasannya tertulis). Perbaiki sampai semua 'OK' / 'OK - new item'.", None),
    ("3. Simpan file, lalu di app: Transaction > Bulk Import > Upload Excel > Read file. App memeriksa lagi; jika semua OK, klik Import.", None),
    ("", None),
    ("Kolom", "h"),
    ("Area (wajib): pilih dari dropdown (area business line Anda).", None),
    ("Item name: wajib. Nama barang adalah identitasnya: barang dicocokkan ke katalog lewat nama, dan dua barang tidak boleh bernama sama.", None),
    ("BPN / SPN: opsional dan boleh sama di beberapa barang. SPN yang sama di dua rak dalam satu area = dua barang dengan nama berbeda (mis. tambahkan nama rak).", None),
    ("Status: diisi otomatis (jangan diketik). 'OK - new item' = barang baru yang akan ditambahkan ke katalog.", None),
    ("Category, UOM: hanya wajib untuk barang BARU (belum ada di katalog).", None),
    ("Qty: > 0 = dicatat sebagai Receive (1 transaksi per area). 0 / kosong = barang hanya ditambahkan ke area.", None),
    ("Min in area: minimum stok di area itu. Kosong = tidak mengubah min yang ada (barang baru: 0).", None),
    ("Note: opsional. Catatan posisi sisa barang di rak lain (mis. 'sisa 2 di Cartridge Filter 2'); kosong = note lama tidak diubah.", None),
    ("Rack: opsional. Dropdown hanya menampilkan rak milik Area di baris itu. Selama Area kosong, sel Rack abu-abu dan dropdown hanya berisi '(Select the Area first)'; sel jadi merah bila Rack diisi tanpa Area.", None),
    ("", None),
    ("Contoh baris", "h"),
    ("Container Filter | Filter Testing | | | | | 5 | 2 | Oil Filter   (barang lama, +5 stok, min 2, rak Oil Filter)", None),
    ("Storage Room | SAFETY GLOVES SIZE L | BPN-GLV-001 | SPN-778 | CNM | PR | 20 | 5 |   (barang baru)", None),
    ("", None),
    ("Catatan", "h"),
    ("Jangan hapus/ganti nama tabel 'tblImport' dan judul kolomnya; jangan ubah sheet 'Lists'.", None),
    ("Nama barang + area yang sama hanya boleh muncul sekali per import (1 rak per barang per area). Barang yang sama di rak lain: beri nama berbeda, atau tulis posisi sisanya di kolom Note.", None),
    ("Download template baru dari app bila ada area/rak/kategori baru.", None),
]
for i, (text, kind) in enumerate(lines, 1):
    cell = gs.cell(row=i, column=2, value=text)
    cell.font = Font(name=F, bold=kind in ("title", "h"), size=14 if kind == "title" else 10 if kind is None else 11,
                     color=NAVY if kind else None)
    cell.alignment = Alignment(wrap_text=True, vertical="top")

wb.active = 1
wb.save(OUT)
print(OUT)
